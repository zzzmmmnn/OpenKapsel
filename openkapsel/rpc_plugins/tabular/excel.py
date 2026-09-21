"""Read-only, bounded Excel adapters. Never save, calculate or execute macros."""
from __future__ import annotations

import importlib.util
import time
import zipfile

from .._data import fail, identity, json_view, object_schema, validate
from .csv_stream import MAX_COLUMNS, CURSOR_TTL_SECONDS, _digest, _encode_cursor, _decode_cursor

MAX_EXCEL_BYTES = 64 * 1024 * 1024
MAX_EXCEL_EXPANDED_BYTES = 512 * 1024 * 1024
EXCEL_SCHEMA = object_schema({
    "sheet": {"type": "string", "minLength": 1, "maxLength": 128},
    "header_row": {"type": "integer", "minimum": 0, "maximum": 1000, "default": 1},
    "start_row": {"type": "integer", "minimum": 1, "maximum": 1048576},
    "value_mode": {"type": "string", "enum": ["cached", "formula"], "default": "cached"},
})


def available_formats():
    formats = []
    if importlib.util.find_spec("openpyxl") and importlib.util.find_spec("defusedxml"):
        formats.extend(["xlsx", "xlsm"])
    if importlib.util.find_spec("xlrd"):
        formats.append("xls")
    return formats


def _zip_preflight(stream):
    try:
        with zipfile.ZipFile(stream) as bundle:
            entries = bundle.infolist()
            if len(entries) > 4096 or sum(info.file_size for info in entries) > MAX_EXCEL_EXPANDED_BYTES:
                fail("excel_expansion_limit", "Excel ZIP entry count or expanded size exceeds the limit", 413)
            names = set()
            for info in entries:
                if info.filename in names or info.flag_bits & 1:
                    fail("excel_archive_invalid", "duplicate or encrypted Excel ZIP members are unsupported", 422)
                names.add(info.filename)
                if info.filename == "xl/sharedStrings.xml":
                    maximum = 8 * 1024 * 1024
                elif info.filename.startswith("xl/worksheets/"):
                    maximum = 128 * 1024 * 1024
                elif info.filename.endswith((".xml", ".rels")):
                    maximum = 2 * 1024 * 1024
                else:
                    maximum = 64 * 1024 * 1024
                if info.file_size > maximum:
                    fail("excel_part_limit", "an Excel part exceeds its bounded parser limit", 413)
    except zipfile.BadZipFile:
        fail("excel_invalid", "file is not a valid OOXML workbook", 422)
    stream.seek(0)


class ExcelStream:
    """Worksheet iterator; Excel continuation rescans compressed worksheet XML."""
    def __init__(self, snapshot, fmt, options, *, cursor=None, task=None):
        validate(options, EXCEL_SCHEMA, "excel")
        if fmt not in available_formats():
            fail("tabular_dependency_missing", "install openpyxl + defusedxml for xlsx/xlsm, or xlrd for xls", 415)
        self.snapshot, self.fmt, self.task = snapshot, fmt, task
        self.options = {"header_row": 1, "value_mode": "cached", **options}
        self.book = None
        self.iterator = None
        self.ended = False
        self.scope = _digest(str(snapshot.path))
        self.version = _digest(identity(snapshot.stat))
        self.dialect = _digest([fmt, {k: v for k, v in self.options.items() if k != "start_row"}])
        try:
            if task:
                task.check_cancelled()
            if fmt in {"xlsx", "xlsm"}:
                _zip_preflight(snapshot.stream)
                from openpyxl import load_workbook
                from openpyxl.xml.functions import DEFUSEDXML
                if not DEFUSEDXML:
                    fail("excel_xml_safety", "openpyxl must have defusedxml protection enabled", 415)
                self.book = load_workbook(snapshot.stream, read_only=True, data_only=self.options["value_mode"] == "cached",
                                          keep_links=False, keep_vba=False)
                self.sheets = [{"name": s.title, "reported_rows": s.max_row, "reported_columns": s.max_column}
                               for s in self.book.worksheets]
                sheet_name = self.options.get("sheet", self.book.sheetnames[0] if self.book.sheetnames else "")
                if sheet_name not in self.book.sheetnames:
                    fail("excel_sheet_missing", "requested worksheet does not exist", 404)
                self.sheet = self.book[sheet_name]
                # Dimensions are producer hints, not a trusted end-of-data marker.
                self.sheet.reset_dimensions()
                header_row = self.options["header_row"]
                header_iterator = self.sheet.iter_rows(min_row=max(1, header_row), max_row=max(1, header_row), values_only=True)
                try:
                    first = next(header_iterator, ())
                finally:
                    header_iterator.close()
                count = len(first)
            else:
                if self.options["value_mode"] == "formula":
                    fail("excel_formula_unavailable", "xls exposes cached values, not formula source", 415)
                import xlrd
                raw = snapshot.stream.read(MAX_EXCEL_BYTES + 1)
                if len(raw) > MAX_EXCEL_BYTES:
                    fail("excel_size_limit", "xls exceeds the 64 MiB input limit", 413)
                self.book = xlrd.open_workbook(file_contents=raw, on_demand=True)
                sheet_name = self.options.get("sheet", self.book.sheet_names()[0] if self.book.nsheets else "")
                if sheet_name not in self.book.sheet_names():
                    fail("excel_sheet_missing", "requested worksheet does not exist", 404)
                self.sheet = self.book.sheet_by_name(sheet_name)
                self.sheets = [{"name": name} for name in self.book.sheet_names()]
                header_row = self.options["header_row"]
                first = self._xls_row(max(0, header_row - 1)) if self.sheet.nrows else []
                count = self.sheet.ncols
            if count > MAX_COLUMNS:
                fail("excel_column_limit", "worksheet header exceeds 4096 columns", 413)
            if header_row:
                self.names = ["" if v is None else str(v) for v in first]
            else:
                self.names = [f"column_{i + 1}" for i in range(count)]
            if sum(map(len, self.names)) > 32768:
                fail("excel_header_limit", "worksheet header exceeds 32768 characters", 413)
            self.row = self.options.get("start_row", header_row + 1 if header_row else 1) - 1
            if self.row < header_row:
                fail("excel_start_row", "start_row must follow header_row")
            if cursor is not None:
                if "start_row" in options:
                    fail("excel_cursor_conflict", "use cursor or start_row, not both")
                state = _decode_cursor(cursor)
                if state.get("scope") != self.scope or state.get("dialect") != self.dialect or state.get("kind") != "excel":
                    fail("tabular_cursor_mismatch", "cursor belongs to another workbook, sheet or value mode", 409)
                if state.get("identity") != self.version:
                    fail("data_source_changed", "workbook changed since this cursor was issued", 409)
                self.row = state["row"]
            self._reset_iterator()
        except BaseException:
            self.close()
            raise

    def _xls_row(self, index):
        import xlrd
        if index >= self.sheet.nrows:
            return []
        result = []
        for cell in self.sheet.row(index):
            if cell.ctype == xlrd.XL_CELL_DATE:
                result.append(xlrd.xldate_as_datetime(cell.value, self.book.datemode).isoformat())
            elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                result.append(bool(cell.value))
            elif cell.ctype == xlrd.XL_CELL_ERROR:
                result.append(xlrd.error_text_from_code.get(cell.value, "#ERROR"))
            elif cell.ctype in {xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK}:
                result.append(None)
            else:
                result.append(cell.value)
        return result

    def _reset_iterator(self):
        if self.fmt in {"xlsx", "xlsm"}:
            if self.iterator is not None and hasattr(self.iterator, "close"):
                self.iterator.close()
            self.iterator = iter(self.sheet.iter_rows(min_row=self.row + 1, values_only=True))

    def next_row(self):
        if self.task:
            self.task.check_cancelled()
        if self.ended:
            return None
        if self.fmt in {"xlsx", "xlsm"}:
            try:
                raw = next(self.iterator)
            except StopIteration:
                self.ended = True
                return None
            values, _types = json_view(list(raw))
        else:
            if self.row >= self.sheet.nrows:
                self.ended = True
                return None
            values = self._xls_row(self.row)
        self.row += 1
        if len(values) > MAX_COLUMNS:
            fail("excel_column_limit", "worksheet row exceeds 4096 columns", 413)
        return values

    def checkpoint(self):
        return self.row

    def restore(self, checkpoint):
        self.row, self.ended = checkpoint, False
        self._reset_iterator()

    def at_eof(self):
        return self.ended  # One final empty page can be necessary for Excel.

    def cursor(self):
        return _encode_cursor({"v": 1, "kind": "excel", "scope": self.scope, "identity": self.version,
                               "dialect": self.dialect, "row": self.row,
                               "expires": int(time.time()) + CURSOR_TTL_SECONDS})

    def close(self):
        if self.iterator is not None and hasattr(self.iterator, "close"):
            self.iterator.close()
        if self.book is not None:
            if self.fmt == "xls":
                self.book.release_resources()
            else:
                self.book.close()
            self.book = None

    def __enter__(self):
        return self

    def __exit__(self, *args):
        self.close()
